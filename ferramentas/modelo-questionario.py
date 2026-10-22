# Gera a planilha modelo do Questionário da aula (assets/modelo-questionario.xlsx), que o professor
# baixa na Área do professor, preenche e envia de volta. A 1ª aba é a que o jogo lê.
# Uso: python ferramentas/modelo-questionario.py
from pathlib import Path
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SAIDA = Path(__file__).resolve().parent.parent / 'assets' / 'modelo-questionario.xlsx'
FONTE = 'Arial'
TURQUESA, VERDE, VERMELHO, AMARELO = '1FA2B0', 'E8F6EC', 'FDECEA', 'FFF6C7'
borda = Border(*(Side(style='thin', color='CFE3E7'),) * 4)

wb = Workbook()
ws = wb.active
ws.title = 'Perguntas'
cabecalho = ['Pergunta', 'Resposta certa', 'Resposta errada 1', 'Resposta errada 2', 'Resposta errada 3']
ws.append(cabecalho)
exemplos = [
    ['Qual é o maior planeta do Sistema Solar?', 'Júpiter', 'Marte', 'Terra', 'Saturno'],
    ['Qual palavra está escrita corretamente?', 'exceção', 'excessão', 'esceção', 'exeção'],
]
for linha in exemplos:
    ws.append(linha)

for col, largura in zip('ABCDE', (52, 24, 24, 24, 24)):
    ws.column_dimensions[col].width = largura
for c in ws[1]:
    c.font = Font(name=FONTE, bold=True, color='FFFFFF', size=11)
    c.fill = PatternFill('solid', fgColor=TURQUESA)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = borda
ws.row_dimensions[1].height = 30
# área de preenchimento (até 200 perguntas) com as cores da legenda: certa em verde, erradas em vermelho
for r in range(2, 202):
    for col in range(1, 6):
        c = ws.cell(row=r, column=col)
        c.font = Font(name=FONTE, size=11, italic=r <= 1 + len(exemplos), color='5B6B7B' if r <= 1 + len(exemplos) else '22364A')
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = borda
        if col == 2:
            c.fill = PatternFill('solid', fgColor=VERDE)
        elif col >= 3:
            c.fill = PatternFill('solid', fgColor=VERMELHO)
ws.freeze_panes = 'A2'
ws['A1'].comment = Comment('Uma pergunta por linha. As 2 linhas em cinza são exemplos: apague ou escreva por cima.', 'Jogos Educacionais')
ws['B1'].comment = Comment('Só a resposta certa (obrigatória).', 'Jogos Educacionais')
ws['C1'].comment = Comment('Pelo menos 1 resposta errada; até 3. No jogo, as respostas aparecem embaralhadas.', 'Jogos Educacionais')

# aba de instruções (a 2ª aba; o jogo lê só a 1ª)
ajuda = wb.create_sheet('Como usar')
ajuda.column_dimensions['A'].width = 4
ajuda.column_dimensions['B'].width = 95
linhas = [
    ('Questionário da aula · planilha modelo', 'titulo'),
    ('', None),
    ('Como preencher', 'sub'),
    ('1. Use a aba "Perguntas" (a primeira). Uma pergunta por linha, a partir da linha 2.', None),
    ('2. Coluna A: a pergunta.', None),
    ('3. Coluna B (verde): a resposta certa.', 'verde'),
    ('4. Colunas C, D e E (vermelho): as respostas erradas. Precisa de pelo menos uma; o ideal são três.', 'vermelho'),
    ('5. As 2 linhas em cinza são exemplos: apague ou escreva por cima antes de enviar.', 'amarelo'),
    ('', None),
    ('Como enviar', 'sub'),
    ('1. No jogo, abra a 🧑‍🏫 Área do professor e digite a senha.', None),
    ('2. Em "📝 Questionários da aula", abra "➕ Novo questionário" e dê um nome (ex.: Ciências – Sistema Solar).', None),
    ('3. Em "📄 Enviar planilha", escolha este arquivo. As perguntas aparecem para conferir.', None),
    ('4. Toque em "Salvar questionário". Ele vira uma aba "Q: nome" na planilha do Google.', None),
    ('5. Toque em "🔗 Copiar link" e mande o link para a turma: ele abre direto este questionário.', None),
    ('', None),
    ('Dicas', 'sub'),
    ('• No jogo as respostas são embaralhadas: não importa em que coluna errada cada uma fica.', None),
    ('• Respostas curtas e claras funcionam melhor no celular.', None),
    ('• Também dá para salvar como .csv (separado por ponto e vírgula) e enviar.', None),
]
for i, (texto, estilo) in enumerate(linhas, start=1):
    c = ajuda.cell(row=i, column=2, value=texto)
    c.font = Font(name=FONTE, size=11, color='22364A')
    c.alignment = Alignment(vertical='center', wrap_text=True)
    if estilo == 'titulo':
        c.font = Font(name=FONTE, size=16, bold=True, color=TURQUESA)
        ajuda.row_dimensions[i].height = 26
    elif estilo == 'sub':
        c.font = Font(name=FONTE, size=12, bold=True, color='137784')
    elif estilo in ('verde', 'vermelho', 'amarelo'):
        c.fill = PatternFill('solid', fgColor={'verde': VERDE, 'vermelho': VERMELHO, 'amarelo': AMARELO}[estilo])

wb.active = 0
SAIDA.parent.mkdir(parents=True, exist_ok=True)
wb.save(SAIDA)
print('Planilha modelo gravada em', SAIDA)
